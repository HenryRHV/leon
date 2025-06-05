from pathlib import Path
import subprocess
import json

from bridges.python.src.sdk.leon import leon
from bridges.python.src.sdk.types import ActionParams

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xmindparser
except ImportError:
    xmindparser = None

try:
    from PIL import Image
except ImportError:
    Image = None


def read_txt(path: Path) -> str:
    return path.read_text(errors='ignore')[:200]


def read_pdf(path: Path) -> str:
    if PyPDF2 is None:
        return 'PyPDF2 not installed'
    text = ''
    with open(path, 'rb') as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages[:5]:
            text += page.extract_text() or ''
    return text[:200]


def read_docx(path: Path) -> str:
    if docx is None:
        return 'python-docx not installed'
    doc = docx.Document(path)
    text = '\n'.join(p.text for p in doc.paragraphs)
    return text[:200]


def read_xlsx(path: Path) -> str:
    if openpyxl is None:
        return 'openpyxl not installed'
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(max_row=5, values_only=True):
        rows.append(', '.join(str(c) for c in row))
    return '\n'.join(rows)


def read_eml(path: Path) -> str:
    import email
    from email import policy
    with open(path, 'rb') as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == 'text/plain':
                return part.get_payload(decode=True).decode(errors='ignore')[:200]
    else:
        return msg.get_payload(decode=True).decode(errors='ignore')[:200]
    return ''


def read_mp4(path: Path) -> str:
    result = subprocess.run([
        'ffprobe', '-v', 'quiet', '-print_format', 'json', '-show_format',
        '-show_streams', str(path)
    ], capture_output=True)
    if result.returncode != 0:
        return 'ffprobe error'
    info = json.loads(result.stdout or '{}')
    duration = info.get('format', {}).get('duration')
    return f'duration: {duration}s'


def read_image(path: Path) -> str:
    if Image is None:
        return 'Pillow not installed'
    with Image.open(path) as img:
        return f'resolution: {img.width}x{img.height}'


def run(params: ActionParams) -> None:
    file_path = params.current_entities[0]['resolution']['value'] if params.current_entities else None
    if not file_path:
        leon.answer({'key': 'not_found', 'data': {'file': ''}})
        return
    path = Path(file_path)
    if not path.exists():
        leon.answer({'key': 'not_found', 'data': {'file': file_path}})
        return

    ext = path.suffix.lower()
    info = ''
    if ext == '.txt':
        info = read_txt(path)
    elif ext == '.pdf':
        info = read_pdf(path)
    elif ext == '.docx':
        info = read_docx(path)
    elif ext == '.xlsx':
        info = read_xlsx(path)
    elif ext == '.eml':
        info = read_eml(path)
    elif ext == '.mp4':
        info = read_mp4(path)
    elif ext in ['.png', '.jpg', '.jpeg']:
        info = read_image(path)
    elif ext == '.xmind':
        if xmindparser is None:
            info = 'xmindparser not installed'
        else:
            data = xmindparser.xmind_to_dict(str(path))
            info = json.dumps(data)[:200]
    else:
        leon.answer({'key': 'unsupported', 'data': {'file': file_path}})
        return

    leon.answer({'key': 'summary', 'data': {'file': file_path, 'info': info}})
